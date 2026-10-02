#!/usr/bin/env python3
"""Génère data/ldc-1991-2026.json.gz (Ligue des Champions, saisons 1991-92 à 2025-26) à partir du
classeur « LDC_Joueurs_cartes_1991-92_2023-24.xlsx » (onglets Participations et Clubs-Saisons).

Usage : python3 scripts/build-ldc-data.py LDC_Joueurs_cartes_1991-92_2023-24.xlsx

- Notes et statistiques du classeur importées telles quelles (le jeu n'a pas de plafond : 99). La rareté
  du classeur (Bronze…Icône) est ignorée : l'import applique le barème officiel (lib/rarity.js).
- Gardiens : colonnes PLO, ARR, JAP, RÉF, VIT, PLA rangées dans les colonnes communes (comme les autres
  imports : pace, shooting, passing, dribbling, defending, physical).
- Postes : premier poste listé ; DF -> CB, FW -> ST, MF -> DM / AM / CM selon les statistiques.
- Noms de clubs ramenés à un nom unique (ALIAS) : même club écrit de plusieurs façons selon les saisons,
  ou nom déjà utilisé par les championnats de la base (pour retrouver son logo).
- Nationalité : absente du classeur ; reprise des cartes des championnats (data/europe-2010-2027.json.gz)
  quand le nom du joueur n'y a qu'une seule nationalité, sinon vide.
"""
import gzip
import json
import os
import re
import sys
import unicodedata

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'data', 'ldc-1991-2026.json.gz')
EUROPE = os.path.join(ROOT, 'data', 'europe-2010-2027.json.gz')

ALIAS = {
    'AC Parma': 'Parma Calcio 1913', 'AFC Ajax': 'Ajax', 'Ajax Amsterdam': 'Ajax',
    'Aston Villa F.C.': 'Aston Villa', 'Athletic Club': 'Athletic Bilbao', 'Atletico Madrid': 'Atlético de Madrid',
    'Bayer Leverkusen': 'Bayer 04 Leverkusen', 'Celta Vigo': 'Celta de Vigo',
    'Deportivo La Coruna': 'Deportivo de La Coruña', 'FC Bayern Munich': 'Bayern Munich',
    'FC Kaiserslautern': '1.FC Kaiserslautern', 'Girondins Bordeaux': 'FC Girondins de Bordeaux',
    'Inter Mediolan': 'Inter Milan', 'Leeds United F.C.': 'Leeds United', 'Manchester City F.C.': 'Manchester City',
    'Manchester United F.C.': 'Manchester United', 'Montpellier HSC': 'Montpellier Hérault Sport Club',
    'Newcastle United F.C.': 'Newcastle United', 'Olympique Lyon': 'Olympique Lyonnais',
    'Olympique Marseille': 'Olympique de Marseille', 'Paris Saint-Germain F.C.': 'Paris Saint-Germain',
    'Real Betis': 'Real Betis Balompié', 'Stade Brestois 29': 'Stade Brestois', 'TSG Hoffenheim': 'TSG 1899 Hoffenheim',
    'Werder Bremen': 'SV Werder Bremen',
    # Même club, plusieurs écritures dans le classeur.
    'Club Brugge KV': 'Club Brugge', 'F.C. Copenhagen': 'FC Copenhagen', 'FC Kobenhavn': 'FC Copenhagen',
    'FC Salzburg': 'Red Bull Salzburg', 'Casino Salzburg': 'Red Bull Salzburg', 'PAE Olympiakos SFP': 'Olympiacos SFP',
    'SK Slavia Praha': 'Slavia Prague', 'SK Sturm Graz': 'Sturm Graz', 'Young Boys': 'BSC Young Boys',
}
PATH_RANK = {'Vainqueur': 0, 'Finale perdue': 1, 'Demi-finales': 2, 'Quarts de finale': 3,
             'Huitièmes / 2e phase de groupes': 4, 'Huitièmes / 2e phase': 4, 'Barrages': 5,
             'Phase de ligue': 6, 'Phase de groupes': 6}


def norm(s):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode().lower())


def season_key(s):
    """'1991-92' -> '1991-1992' ; '1999-00' -> '1999-2000'."""
    start = int(s[:4])
    return f'{start}-{start + 1}'


def sub_position(tokens, s):
    first = (tokens or 'MF').split(',')[0].strip()
    if first in ('GK', 'RB', 'LB', 'CB', 'DM', 'CM', 'AM', 'RW', 'ST', 'LW'):
        return first
    if first == 'DF':
        return 'CB'
    if first == 'FW':
        return 'ST'
    pace, shoot, pas, dri, de, phy = s
    return 'DM' if de - shoot >= 8 else 'AM' if shoot - de >= 8 else 'CM'


def stats(row6, position_tokens):
    c = [int(v) for v in row6]
    if (position_tokens or '').split(',')[0].strip() == 'GK':   # PLO, ARR, JAP, RÉF, VIT, PLA
        c = [c[4], c[0], c[1], c[2], c[3], c[5]]
    return c


def european_nations():
    if not os.path.exists(EUROPE):
        return {}
    d = json.load(gzip.open(EUROPE))
    seen = {}
    for r in d['rows']:
        seen.setdefault(norm(r[3]), set()).add(d['nations'][r[4]])
    return {k: next(iter(v)) for k, v in seen.items() if len(v) == 1}


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
    nations = european_nations()
    clubs, rows, path = [], [], {}

    def club_index(name):
        name = ALIAS.get(name, name)
        if name not in clubs:
            clubs.append(name)
        return clubs.index(name)

    for r in wb['Clubs-Saisons'].iter_rows(min_row=2, values_only=True):
        if r[0]:
            path[(season_key(r[0]), club_index(r[1]))] = PATH_RANK[r[2]]
    seen = {}
    for r in wb['Participations'].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        season, name, pos, club, pid, ovr = season_key(r[0]), str(r[1]).strip(), r[2], club_index(r[3]), r[4], int(r[13])
        s = stats(r[14:20], pos)
        if not name or any(not 1 <= v <= 99 for v in s + [ovr]):
            raise SystemExit(f'Ligne invalide : {r[:5]}')
        key = (season, club, name, pid or '')
        seen[key] = seen.get(key, 0) + 1          # homonymes / doublons : compteur dans l'identifiant
        rows.append([season, club, name, sub_position(pos, s), ovr] + s + [pid or 0, seen[key], nations.get(norm(name), '')])
        path.setdefault((season, club), PATH_RANK['Phase de groupes'])
    seasons = {}
    for (season, club), rank in sorted(path.items(), key=lambda kv: (kv[0][0], kv[1], clubs[kv[0][1]])):
        seasons.setdefault(season, []).append(club)
    doc = {
        'source': os.path.basename(sys.argv[1]),
        'clubs': clubs,
        'seasons': [[s, c] for s, c in sorted(seasons.items())],
        'columns': ['season', 'club', 'name', 'position', 'ovr', 'pace', 'shooting', 'passing', 'dribbling',
                    'defending', 'physical', 'playerId', 'n', 'nation'],
        'rows': rows,
    }
    with gzip.GzipFile(OUT, 'wb', mtime=0) as f:
        f.write(json.dumps(doc, ensure_ascii=False, separators=(',', ':')).encode('utf-8'))
    with_nation = sum(1 for r in rows if r[-1])
    print(f'{len(rows)} cartes, {len(clubs)} clubs, {len(seasons)} saisons, note max {max(r[4] for r in rows)}, '
          f'nationalité retrouvée {with_nation}/{len(rows)}')


if __name__ == '__main__':
    main()
