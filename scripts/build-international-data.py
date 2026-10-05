#!/usr/bin/env python3
"""Génère data/international-2026.json.gz (Coupe du monde 1930-2026 et Ligue des Nations 2019-2025)
à partir des deux classeurs fournis, en appliquant les règles de correction des notes « hors sol ».

Usage :
  python3 scripts/build-international-data.py coupes_du_monde_effectifs_notes.xlsx ligue_des_nations_effectifs_notes.xlsx

Notes : Ligue des Nations importée telle quelle (le jeu n'a pas de plafond : note maximale 99).
Coupe du monde : les notes du classeur sont trop basses (médiane 66, 89 % de Common) ; elles sont relevées
par la courbe LIFT ci-dessous (croissante : l'ordre des joueurs est conservé, les notes >= 92 ne bougent
pas, donc Messi, Pelé, Mbappé... restent à 98). Les six statistiques sont mises à l'échelle (nouvelle note /
ancienne) pour rester cohérentes avec la note affichée.
Nettoyages : âge négatif -> valeur absolue ; numéro ou âge absents -> vides ; gardiens rangés dans les
colonnes communes ; postes DF/MF/FW -> CB/DM-CM-AM/ST
(les côtés RB/LB/RW/LW sont ajoutés ensuite par scripts/refine-international-positions.js).
Le fichier JSON est lu par scripts/import-international.js (identifiants stables, import sans doublon).
"""
import gzip
import json
import os
import sys

import openpyxl

# Relèvement des notes de la Coupe du monde (points d'ancrage note du classeur -> note retenue).
LIFT = [(54, 64), (66, 75), (78, 84), (88, 90), (92, 92)]


def lift(ovr):
    """Interpolation linéaire entre les points d'ancrage ; inchangé au-delà de 92."""
    if ovr >= LIFT[-1][0]:
        return ovr
    if ovr <= LIFT[0][0]:
        return ovr + LIFT[0][1] - LIFT[0][0]
    for (x0, y0), (x1, y1) in zip(LIFT, LIFT[1:]):
        if x0 <= ovr <= x1:
            return round(y0 + (y1 - y0) * (ovr - x0) / (x1 - x0))


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'data', 'international-2026.json.gz')


def stats(row6, position, factor):
    c = [int(v) for v in row6]
    if position == 'GK':            # PLO, ARR, JAP, RÉF, VIT, PLA -> colonnes communes comme les 4 championnats
        c = [c[4], c[0], c[1], c[2], c[3], c[5]]  # pace, shooting, passing, dribbling, defending, physical
    return [max(1, min(99, round(v * factor))) for v in c]


def sub_position(position, s):
    """DF -> CB ; MF -> DM / CM / AM selon les stats ; FW -> ST (les côtés ne sont pas dans les sources)."""
    if position == 'GK':
        return 'GK'
    pace, shoot, pas, dri, de, phy = s
    if position == 'DF':
        return 'CB'
    if position == 'MF':
        if de - shoot >= 8:
            return 'DM'
        if shoot - de >= 8:
            return 'AM'
        return 'CM'
    return 'ST'


def build(cdm_path, nl_path):
    teams, rows = [], []

    def team_index(name):
        if name not in teams:
            teams.append(name)
        return teams.index(name)

    wb = openpyxl.load_workbook(cdm_path, read_only=True, data_only=True)
    # 1930-2022 : Édition, Sélection, N°, Joueur, Poste, Âge, Matchs, Buts, Général, 6 stats, Rareté
    for r in wb['Hommes 1930-2022'].iter_rows(min_row=2, values_only=True):
        edition, team, number, name, pos, age, ovr = r[0], r[1], r[2], r[3], r[4], r[5], r[8]
        new = lift(int(ovr))
        s = stats(r[9:15], pos, new / int(ovr))
        s0 = stats(r[9:15], pos, 1)   # poste déduit des statistiques d'origine (le relèvement ne change pas le poste)
        age = abs(int(age)) if age is not None else None
        rows.append(['coupe-du-monde', str(edition), team_index(team), name, sub_position(pos, s0), number, age, new] + s)
    # 2026 : Édition, Sélection, N°, Joueur, Poste, Âge, Sél., Buts int., Club, Général, 6 stats, Rareté
    for r in wb['Coupe du monde 2026'].iter_rows(min_row=2, values_only=True):
        edition, team, number, name, pos, age, ovr = r[0], r[1], r[2], r[3], r[4], r[5], r[9]
        new = lift(int(ovr))
        s = stats(r[10:16], pos, new / int(ovr))
        s0 = stats(r[10:16], pos, 1)
        rows.append(['coupe-du-monde', str(edition), team_index(team), name, sub_position(pos, s0), number,
                     int(age) if age is not None else None, new] + s)

    wl = openpyxl.load_workbook(nl_path, read_only=True, data_only=True)
    season = {2019: '2018-2019', 2021: '2020-2021', 2023: '2022-2023', 2025: '2024-2025'}
    for r in wl['Phases finales 2019-2025'].iter_rows(min_row=2, values_only=True):
        edition, team, number, name, pos, age, ovr = r[0], r[1], r[2], r[3], r[4], r[5], r[9]
        new = int(ovr)
        s = stats(r[10:16], pos, 1)
        rows.append(['ligue-des-nations', season[int(edition)], team_index(team), name, sub_position(pos, s), number,
                     int(age) if age is not None else None, new] + s)
    return teams, rows


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    teams, rows = build(sys.argv[1], sys.argv[2])
    doc = {
        'source': 'coupes_du_monde_effectifs_notes.xlsx + ligue_des_nations_effectifs_notes.xlsx',
        'competitions': ['coupe-du-monde', 'ligue-des-nations'],
        'teams': teams,
        'columns': ['competition', 'season', 'team', 'name', 'position', 'number', 'age', 'ovr', 'pace', 'shooting',
                    'passing', 'dribbling', 'defending', 'physical'],
        'rows': rows,
    }
    with gzip.GzipFile(OUT, 'wb', mtime=0) as f:   # mtime=0 : fichier reproductible (même contrôle d'intégrité)
        f.write(json.dumps(doc, ensure_ascii=False, separators=(',', ':')).encode('utf-8'))
    print(f'{len(rows)} cartes, {len(teams)} sélections, note max {max(r[7] for r in rows)}')


if __name__ == '__main__':
    main()
